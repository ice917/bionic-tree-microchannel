# -*- coding: utf-8 -*-
"""数据链路口径断言（阶段 0）。

把本次审计暴露的「口径锚点」固化成回归断言，防止数据链路再次断裂。
每个断言对应一个曾被静默破坏、或一旦破坏结论就变的口径：

  A1  单源压降比 310×（run_report.json，约定 A 网络公式）
  A2  多源等钉住 dp 比最大值 96.9×（multi_source_summary.json，网络公式；
      4 芯片/1300 格。平衡分配 + 实测半径修复前为 49.5×）
  A3  verify.poiseuille_dp 默认走网络路径求解（as_network=True），
      网络值必须远小于串行求和值
  A4  calibrate.tree_dp_pa 返回网络 dp（~0.0802 kPa），
      串行参考值 ~6.42 kPa（80× 过计），两参考字段必须都在场
  A5  field.solve_heat_sources 的 q 无条件缩放（q=2 的温升是 q=1 的 2 倍）
  A6  单源树与直通道峰值温度逐位相同（27.596 °C）
  A7  三维共轭的外迭代全局上界 155、末值阻尼不是固定 0.5（conjugate_3d_flow.json）
  A8  三份稿件不得出现已被代码证伪的口径（固定松弛 / 「每次均 122 次」/
      按下游格数分流 / 「currently v1.0.x」时效枚举）；IJTS 稿更进一步，
      全文不得出现任何 vN.N.N 字面（版本枚举在下次归档即失效）
  A9  三份稿件必须出现已订正口径（下游叶数分流 / 155 次上界）

A7–A9 的存在理由：稿件里的这三句一度只在中英文各一份里被订正，IJTS 稿滞后到
v1.0.7 仍保留旧口径。把它们与 outputs 直接绑定，改稿不同步即 FAIL。

退出码 0 = 全部 PASS；非 0 = 有 FAIL。挂到论文生成链路前执行。
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"

FAIL = []
PASS = 0


def _load(name):
    with open(OUT / name, encoding="utf-8") as f:
        return json.load(f)


def check(name, cond, detail=""):
    global PASS
    if cond:
        PASS += 1
        print("  [OK]   %s %s" % (name, detail))
    else:
        FAIL.append(name)
        print("  [FAIL] %s %s" % (name, detail))


def approx(got, exp, tol):
    return abs(float(got) - float(exp)) <= tol


# ------------------------------------------------------------------ A1
run = _load("run_report.json")["compare"]
check("A1 单源压降比 310×",
      approx(run["dp_ratio_tree_over_straight"], 310.17, 0.5),
      "= %.2f" % run["dp_ratio_tree_over_straight"])
check("A1 单源树压降 7.3472",
      approx(run["tree"]["pressure_drop"], 7.3472, 1e-3))
check("A1 单源直通道压降 0.0237",
      approx(run["straight"]["pressure_drop"], 0.0237, 1e-4))

# ------------------------------------------------------------------ A2
ms = _load("multi_source_summary.json")
rows = [r for r in ms["rows"] if "tree_cells" in r and "| budget" in r["case"]]
dpmax = max(r["dp_ratio_tree_over_ctrl"] for r in rows)
check("A2 多源等钉住 dp 比最大值 96.9×", approx(dpmax, 96.9, 0.1),
      "= %.1f (修复前为 49.5×)" % dpmax)

# ------------------------------------------------------------------ A3
import grow as G
import radius as Ra
import verify as V
import field as F

T = F.solve_heat_sources(grid=30).T
edges = G.grow_tree(T, max_nodes=120)
radii = Ra.assign_radii(edges, grid=30, r0=0.5)
dp_net = V.poiseuille_dp(edges, radii, 30)                      # 网络
dp_ser = V.poiseuille_dp(edges, radii, 30, as_network=False)    # 串行
check("A3 poiseuille_dp 网络 << 串行", dp_net < dp_ser / 10.0,
      "net=%.4f serial=%.4f" % (dp_net, dp_ser))

# ------------------------------------------------------------------ A4
import calibrate as C

dp_pa = C.tree_dp_pa(edges, radii, 30, C.paper_grid_mm(30), 0.7282e-3 / 60.0)
check("A4 tree_dp_pa 网络 dp 在场", "dp_pa" in dp_pa)
check("A4 串行参考字段在场",
      "dp_serial_reference_pa" in dp_pa and "dp_kirchhoff_reference_pa" in dp_pa,
      "keys=%s" % sorted(dp_pa.keys()))

# ------------------------------------------------------------------ A5
r1 = F.solve_heat_sources(grid=20, q=1.0)
r2 = F.solve_heat_sources(grid=20, q=2.0)
rise1 = float(r1.T.max() - 25.0)
rise2 = float(r2.T.max() - 25.0)
check("A5 q 无条件缩放", approx(rise2 / rise1, 2.0, 1e-4),
      "rise(q=2)/rise(q=1)=%.6f" % (rise2 / rise1))

# ------------------------------------------------------------------ A6
check("A6 单源树==直通道 T_max",
      approx(run["tree"]["T_max_C"], run["straight"]["T_max_C"], 1e-9)
      and approx(run["tree"]["T_max_C"], 27.596, 1e-3),
      "= %.3f °C" % run["tree"]["T_max_C"])

# ------------------------------------------------------------------ A7
# 三维共轭：稿件引用的两个结论性数字必须与 outputs 一致
def _collect(node, key, out):
    if isinstance(node, dict):
        for k, v in node.items():
            if k == key:
                out.append(v)
            else:
                _collect(v, key, out)
    elif isinstance(node, list):
        for v in node:
            _collect(v, key, out)
    return out


c3d = _load("conjugate_3d_flow.json")
iters = _collect(c3d, "picard_iters", [])
damps = _collect(c3d, "picard_damping_final", [])
check("A7 外迭代全局上界 155", max(iters) == 155,
      "max=%d n=%d" % (max(iters), len(iters)))
check("A7 末值阻尼自适应（非固定 0.5）",
      len(set(damps)) >= 2 and min(damps) < 0.5,
      "min=%.4f distinct=%d" % (min(damps), len(set(damps))))

# ------------------------------------------------------------------ A8/A9
# 稿件口径守卫：单链 5 处缺陷（D-1/D-2/D-3）的复发拦截
import docx

MANUSCRIPTS = [
    ("IJTS", ROOT / "论文/submission_ijts/Manuscript_IJTS.docx"),
    ("CTS", ROOT / "论文/submission_cts/Manuscript_CTS_blinded.docx"),
    ("CN", ROOT / "论文/仿生自适应网络芯片散热与布线.docx"),
]

FORBIDDEN = [
    "number of cells in the downstream",      # D-1 错误口径
    "(relaxation 0.5",                        # D-2 把自适应阻尼写成固定值
    "within 122 outer iterations on every run",  # D-2 与 155 次实测冲突
    "按下游子树的格数",                        # D-1 中文旧口径
    "currently v1.0.",                        # D-3 时效枚举
]

REQUIRED = {
    "IJTS": ["number of leaves in the downstream",
             "within 155 outer iterations"],
    "CTS": ["number of leaves in the downstream",
            "within 155 outer iterations"],
    "CN": ["下游子树的叶节点数", "155 次"],
}

for tag, path in MANUSCRIPTS:
    check("A8 %s 稿件存在" % tag, path.exists(), str(path))
    if not path.exists():
        continue
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    for s in FORBIDDEN:
        check("A8 %s 无 %r" % (tag, s), s not in text)
    for s in REQUIRED[tag]:
        check("A9 %s 有 %r" % (tag, s), s in text)

# A8 补充：任何「vN.N.N」字面都是时效性真理断言，下一个版本归档那一刻即可能变假，
# 这正是 D-3 的成因。IJTS 稿已改为「指向 Zenodo 记录页」的不去时效表述，故全文
# 零版本号字面可以严格断言。中文稿的「包括…不再逐版更新本段」是非穷尽 + 自陈
# 冻结，不构成时效断言，因此本条只对 IJTS 生效（不误报）。
_ijts = ROOT / "论文/submission_ijts/Manuscript_IJTS.docx"
if _ijts.exists():
    _ijts_text = "\n".join(p.text for p in docx.Document(str(_ijts)).paragraphs)
    _tokens = sorted(set(re.findall(r"v\d+\.\d+\.\d+", _ijts_text)))
    check("A8 IJTS 零版本号字面（去时效）", not _tokens, "tokens=%s" % _tokens)

# ------------------------------------------------------------------ 结果
print("\n===== %d PASS / %d FAIL =====" % (PASS, len(FAIL)))
if FAIL:
    for f in FAIL:
        print("  FAILED:", f)
    sys.exit(1)
